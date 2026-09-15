import io
import logging
import zipfile
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {".txt", ".md", ".log", ".py", ".java", ".js", ".ts", ".jsx", ".tsx"}


def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_parts: list[str] = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = DocxDocument(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs).strip()


def extract_text_from_plain(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def extract_text_from_zip(file_bytes: bytes, max_files: int = 200, max_file_size: int = 300_000) -> str:
    """Concatenate readable source/text files from a zipped project, each labeled by path."""
    parts: list[str] = []
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        infos = [
            info for info in zf.infolist()
            if not info.is_dir()
            and Path(info.filename).suffix.lower() in TEXT_EXTENSIONS
            and info.file_size <= max_file_size
        ][:max_files]

        for info in infos:
            try:
                content = zf.read(info).decode("utf-8", errors="ignore")
            except Exception:
                continue
            parts.append(f"\n\n===== FILE: {info.filename} =====\n{content}")

    return "".join(parts).strip()


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Dispatch to the correct extractor based on file extension."""
    ext = Path(filename).suffix.lower()

    try:
        if ext == ".pdf":
            return extract_text_from_pdf(file_bytes)
        if ext == ".docx":
            return extract_text_from_docx(file_bytes)
        if ext == ".zip":
            return extract_text_from_zip(file_bytes)
        if ext in TEXT_EXTENSIONS:
            return extract_text_from_plain(file_bytes)
        # Fallback: best-effort plain text decode
        return extract_text_from_plain(file_bytes)
    except Exception:
        logger.exception("Failed to extract text from %s", filename)
        return ""
